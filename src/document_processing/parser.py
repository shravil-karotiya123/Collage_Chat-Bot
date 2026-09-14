"""
Document Parser Module for MRPL AI Workbench.
Extracts raw text content and structure from PDF, DOCX, TXT, MD, CSV, and XLSX formats.
"""

import csv
import io
import re
import zipfile
import xml.etree.ElementTree as ET
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from typing import Any, Dict, List

import docx
import pypdf

from src.document_processing.loader import LoadedDocument


@dataclass
class ParsedDocument:
    """
    Container encapsulating extracted text content and parser structural metadata.
    """

    text_content: str
    char_count: int
    word_count: int
    page_or_sheet_count: int
    metadata: Dict[str, Any] = field(default_factory=dict)


class BaseFormatParser(ABC):
    """
    Abstract Base Class contract for format-specific document parsers.
    """

    @abstractmethod
    def parse(self, doc: LoadedDocument) -> ParsedDocument:
        """Parse raw LoadedDocument bytes into ParsedDocument."""
        pass


class PDFFormatParser(BaseFormatParser):
    """
    PDF text extraction parser using pypdf.
    """

    def parse(self, doc: LoadedDocument) -> ParsedDocument:
        stream = io.BytesIO(doc.content_bytes)
        reader = pypdf.PdfReader(stream)
        page_texts: List[str] = []

        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                page_texts.append(f"--- Page {idx + 1} ---\n{page_text.strip()}")

        full_text = "\n\n".join(page_texts) if page_texts else ""
        char_count = len(full_text)
        word_count = len(full_text.split()) if full_text else 0

        return ParsedDocument(
            text_content=full_text,
            char_count=char_count,
            word_count=word_count,
            page_or_sheet_count=len(reader.pages),
            metadata={"parser": "pypdf", "page_count": len(reader.pages)},
        )


class DOCXFormatParser(BaseFormatParser):
    """
    Word document parser extracting paragraphs and table contents using python-docx.
    """

    def parse(self, doc: LoadedDocument) -> ParsedDocument:
        stream = io.BytesIO(doc.content_bytes)
        docx_doc = docx.Document(stream)
        lines: List[str] = []

        # Extract Paragraphs
        for p in docx_doc.paragraphs:
            if p.text.strip():
                lines.append(p.text.strip())

        # Extract Tables
        for table_idx, table in enumerate(docx_doc.tables):
            lines.append(f"\n--- Table {table_idx + 1} ---")
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                lines.append(" | ".join(row_cells))

        full_text = "\n".join(lines)
        char_count = len(full_text)
        word_count = len(full_text.split()) if full_text else 0

        return ParsedDocument(
            text_content=full_text,
            char_count=char_count,
            word_count=word_count,
            page_or_sheet_count=1,
            metadata={
                "parser": "python-docx",
                "paragraph_count": len(docx_doc.paragraphs),
                "table_count": len(docx_doc.tables),
            },
        )


class TextFormatParser(BaseFormatParser):
    """
    Plain Text (.txt) and Markdown (.md) parser.
    """

    def parse(self, doc: LoadedDocument) -> ParsedDocument:
        try:
            full_text = doc.content_bytes.decode("utf-8")
        except UnicodeDecodeError:
            full_text = doc.content_bytes.decode("latin-1", errors="replace")

        full_text = full_text.strip()
        char_count = len(full_text)
        word_count = len(full_text.split()) if full_text else 0

        return ParsedDocument(
            text_content=full_text,
            char_count=char_count,
            word_count=word_count,
            page_or_sheet_count=1,
            metadata={"parser": "text_decoder", "format": doc.file_extension},
        )


class CSVFormatParser(BaseFormatParser):
    """
    Comma Separated Values (.csv) parser formatting rows into clean structured text.
    """

    def parse(self, doc: LoadedDocument) -> ParsedDocument:
        try:
            raw_str = doc.content_bytes.decode("utf-8")
        except UnicodeDecodeError:
            raw_str = doc.content_bytes.decode("latin-1", errors="replace")

        csv_file = io.StringIO(raw_str)
        reader = csv.reader(csv_file)
        rows = list(reader)

        formatted_lines: List[str] = []
        for idx, row in enumerate(rows):
            if not row:
                continue
            if idx == 0:
                formatted_lines.append("Headers: " + " | ".join(row))
            else:
                formatted_lines.append(f"Row {idx}: " + " | ".join(row))

        full_text = "\n".join(formatted_lines)
        char_count = len(full_text)
        word_count = len(full_text.split()) if full_text else 0

        return ParsedDocument(
            text_content=full_text,
            char_count=char_count,
            word_count=word_count,
            page_or_sheet_count=1,
            metadata={"parser": "csv_reader", "row_count": len(rows)},
        )


class XLSXFormatParser(BaseFormatParser):
    """
    Excel spreadsheet (.xlsx) parser extracting worksheet text using standard zipfile/XML parsing.
    """

    def parse(self, doc: LoadedDocument) -> ParsedDocument:
        stream = io.BytesIO(doc.content_bytes)
        extracted_text_blocks: List[str] = []
        sheet_count = 0

        try:
            with zipfile.ZipFile(stream, "r") as z:
                # 1. Parse Shared Strings
                shared_strings: List[str] = []
                if "xl/sharedStrings.xml" in z.namelist():
                    with z.open("xl/sharedStrings.xml") as f:
                        tree = ET.parse(f)
                        for elem in tree.iter():
                            if elem.tag.endswith("t") and elem.text:
                                shared_strings.append(elem.text.strip())

                # 2. Find and Parse Worksheet Files
                sheet_files = [f for f in z.namelist() if f.startswith("xl/worksheets/sheet") and f.endswith(".xml")]
                sheet_count = len(sheet_files)

                if shared_strings:
                    extracted_text_blocks.append("Shared Strings:\n" + "\n".join(shared_strings))

                # Parse Sheet cell text fallback
                for sheet_file in sheet_files:
                    sheet_name = sheet_file.split("/")[-1].replace(".xml", "")
                    with z.open(sheet_file) as f:
                        tree = ET.parse(f)
                        cell_texts = []
                        for elem in tree.iter():
                            if elem.tag.endswith("v") and elem.text:
                                cell_texts.append(elem.text.strip())
                        if cell_texts:
                            extracted_text_blocks.append(f"--- Sheet: {sheet_name} ---\n" + ", ".join(cell_texts))

        except Exception as exc:
            # Fallback regex extraction from raw bytes
            strings = re.findall(rb"[\x20-\x7E]{4,}", doc.content_bytes)
            decoded = [s.decode("ascii", errors="ignore") for s in strings if len(s) > 4]
            extracted_text_blocks.append("\n".join(decoded))

        full_text = "\n\n".join(extracted_text_blocks)
        char_count = len(full_text)
        word_count = len(full_text.split()) if full_text else 0

        return ParsedDocument(
            text_content=full_text,
            char_count=char_count,
            word_count=word_count,
            page_or_sheet_count=sheet_count or 1,
            metadata={"parser": "xlsx_zip_xml", "sheet_count": sheet_count},
        )


class DocumentParser:
    """
    Main Document Parser routing LoadedDocument to the appropriate format-specific parser.
    """

    def __init__(self) -> None:
        self._parsers: Dict[str, BaseFormatParser] = {
            ".pdf": PDFFormatParser(),
            ".docx": DOCXFormatParser(),
            ".txt": TextFormatParser(),
            ".md": TextFormatParser(),
            ".csv": CSVFormatParser(),
            ".xlsx": XLSXFormatParser(),
        }

    def parse(self, doc: LoadedDocument) -> ParsedDocument:
        """
        Parse LoadedDocument into structured ParsedDocument payload.

        Args:
            doc: LoadedDocument instance.

        Returns:
            ParsedDocument containing extracted text and structural metadata.
        """
        parser = self._parsers.get(doc.file_extension)
        if not parser:
            raise ValueError(f"No parser registered for file extension '{doc.file_extension}'.")

        try:
            return parser.parse(doc)
        except Exception as exc:
            raise RuntimeError(f"Error parsing document '{doc.file_name}': {str(exc)}") from exc
